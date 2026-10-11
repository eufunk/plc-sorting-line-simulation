"""Erzeugt aus src/HMI/TextLists.csv eine Excel-Datei, die TIA Portal unter
HMI_SortingLine -> Text- und Grafiklisten -> Textlisten -> Importieren einlesen kann.

Format nach einem Export aus TIA Portal V21 (Blätter "TextList" und "TextListEntry").
Bereich je Liste: "Decimal" (Wert/Bereich) oder "Bit" (für Bool-Variablen).

Aufruf:  python tools/make_textlist_import.py [Zieldatei]
Standard-Ziel: C:\\TIA\\Textlisten_Import.xlsx (außerhalb des Repositorys)
"""

import csv
import sys
from pathlib import Path

import openpyxl
from openpyxl.packaging.custom import StringProperty

REPO = Path(__file__).resolve().parent.parent
TEXTLISTS_CSV = REPO / "src" / "HMI" / "TextLists.csv"
DEFAULT_TARGET = Path(r"C:\TIA\Textlisten_Import.xlsx")
LANGUAGE = "de-DE"


def main():
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_TARGET
    with open(TEXTLISTS_CSV, encoding="utf-8", newline="") as f:
        rows = list(csv.DictReader(f, delimiter=";"))

    wb = openpyxl.Workbook()

    lists = wb.active
    lists.title = "TextList"
    lists.append(["Name", "ListRange", f"Comment [{LANGUAGE}]"])
    seen = []
    for row in rows:
        if row["Liste"] not in seen:
            seen.append(row["Liste"])
            lists.append([row["Liste"], row["Bereich"], "<No value>"])

    entries = wb.create_sheet("TextListEntry")
    entries.append(["Name", "Parent", "DefaultEntry", "Value", f"Text [{LANGUAGE}]", "FieldInfos"])
    count = {}
    for row in rows:
        n = count[row["Liste"]] = count.get(row["Liste"], 0) + 1
        entries.append([f"{row['Liste']}_{n}", row["Liste"], "", row["Wert"], row["Text"], ""])

    # Alle Zellen als Text, wie im TIA-Export
    for ws in wb.worksheets:
        for r in ws.iter_rows():
            for cell in r:
                cell.number_format = "@"

    # Versionskennung wie im TIA-Export, sonst warnt TIA beim Import
    wb.custom_doc_props.append(StringProperty(name="TIA_Version", value="1.1"))

    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)
    print(f"{len(seen)} Textlisten, {len(rows)} Einträge -> {target}")


if __name__ == "__main__":
    main()
