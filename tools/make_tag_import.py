"""Erzeugt aus src/Constants/Constants.csv und src/Tags/IoTags.csv eine Excel-Datei,
die TIA Portal unter PLC-Variablen -> Variablentabelle -> Importieren einlesen kann.

Format nach einem Export aus TIA Portal V21 (Blätter "PLC Tags", "Constants",
"TagTable Properties"). Die Spalte "Path" bestimmt die Variablentabelle:
  - Konstanten  -> Tabelle "Constants"
  - I/O-Variablen -> Tabelle "IO"

Aufruf:  python tools/make_tag_import.py [Zieldatei]
Standard-Ziel: C:\\TIA\\PlcTags_Import.xlsx (außerhalb des Repositorys)
"""

import csv
import sys
from pathlib import Path

import openpyxl
from openpyxl.packaging.custom import StringProperty

REPO = Path(__file__).resolve().parent.parent
CONSTANTS_CSV = REPO / "src" / "Constants" / "Constants.csv"
IO_TAGS_CSV = REPO / "src" / "Tags" / "IoTags.csv"
DEFAULT_TARGET = Path(r"C:\TIA\PlcTags_Import.xlsx")

CONSTANTS_TABLE = "Constants"
IO_TABLE = "IO"


def read_csv(path):
    with open(path, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f, delimiter=";"))


def main():
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_TARGET
    constants = read_csv(CONSTANTS_CSV)
    io_tags = read_csv(IO_TAGS_CSV)

    wb = openpyxl.Workbook()

    tags = wb.active
    tags.title = "PLC Tags"
    tags.append(["Name", "Path", "Data Type", "Logical Address", "Comment",
                 "Hmi Visible", "Hmi Accessible", "Hmi Writeable",
                 "Typeobject ID", "Version ID"])
    for row in io_tags:
        tags.append([row["Name"], IO_TABLE, row["Datentyp"], row["Adresse"], row["Kommentar"],
                     "True", "True", "True", "", ""])

    consts = wb.create_sheet("Constants")
    consts.append(["Name", "Path", "Data Type", "Value", "Comment"])
    for row in constants:
        consts.append([row["Name"], CONSTANTS_TABLE, row["Datentyp"], row["Wert"], row["Kommentar"]])

    props = wb.create_sheet("TagTable Properties")
    props.append(["Path", "BelongsToUnit", "Accessibility"])
    props.append([CONSTANTS_TABLE, "", ""])
    props.append([IO_TABLE, "", ""])

    # Alle Zellen als Text, wie im TIA-Export (sonst macht Excel aus "0.5" eine Zahl)
    for ws in wb.worksheets:
        for r in ws.iter_rows():
            for cell in r:
                cell.number_format = "@"

    # Versionskennung wie im TIA-Export, sonst warnt TIA beim Import
    wb.custom_doc_props.append(StringProperty(name="TIA_Version", value="1.0"))

    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)
    print(f"{len(constants)} Konstanten, {len(io_tags)} I/O-Variablen -> {target}")


if __name__ == "__main__":
    main()
